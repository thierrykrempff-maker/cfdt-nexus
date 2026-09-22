#!/usr/bin/env python
"""
Nexus local interface server.

Serves a local-only UI and calls the existing Assistant DS Router CLI:
automation/scripts/assistant_ds_router.py ask --query ... --format json
"""

from __future__ import annotations

import argparse
import base64
import importlib.util
import io
import json
import os
import re
import socket
import subprocess
import sys
import unicodedata
import webbrowser
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse


APP_DIR = Path(__file__).resolve().parent
ROOT = APP_DIR.parents[1]
ROUTER_SCRIPT = ROOT / "automation" / "scripts" / "assistant_ds_router.py"
EXPERTS_DIR = ROOT / "automation"
INTERNAL_ERROR_MESSAGE = "Une erreur interne est survenue. Consultez les journaux du serveur."

sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(EXPERTS_DIR))
from experts import orchestrator, report_generator  # noqa: E402
from employee_case_demo import build_demo_payload, public_scenarios  # noqa: E402
from historical_cases import (  # noqa: E402
    get_historical_case,
    list_historical_cases,
    refresh_historical_case_sources,
)
from local_cases import LocalCaseStore  # noqa: E402
from local_ocr import LocalOCRUnavailable, extract_text_from_jpeg  # noqa: E402
from cse_agenda_analysis import build_cse_agenda_analysis  # noqa: E402
from rgpd_analysis import build_rgpd_analysis, build_rgpd_corpus_scan  # noqa: E402
from NEXUS_RUNTIME_INTEGRATION import (  # noqa: E402
    RuntimeCoreIntegration,
    RuntimeCoreIntegrationInput,
    RuntimeCoreReportMapper,
    RuntimeConnectorConfig,
    RuntimeConnectorPayloadMapper,
    RuntimeCSEMemoryConfig,
    RuntimeCSEMemoryDiagnostics,
    RuntimeCSEMemoryIntegration,
    RuntimeCSEMemoryMode,
    RuntimeCSEMemoryReportMapper,
    RuntimeCSEMemoryResult,
    RuntimeIntegrationConfig,
    RuntimeFinalAssistantConfig,
    RuntimeOfficialConnectorsConfig,
    RuntimeOfficialConnectorsIntegration,
    RuntimeProtectionSocialeConfig,
    RuntimeProtectionSocialeIntegration,
    RuntimeProtectionSocialeReportMapper,
    RuntimeRetirementConfig,
    RuntimeRetirementIntegration,
    RuntimeRetirementReportMapper,
    RuntimeSyndicalReasoningConfig,
    RuntimeSyndicalReasoningIntegration,
    RuntimeSyndicalReasoningReportMapper,
    merge_metadata_source_qualifications,
    get_nexus_version,
    sanitize_public_payload,
)
from NEXUS_RUNTIME_INTEGRATION.public_payload import (  # noqa: E402
    sanitize_http_public_payload,
)


OPTIONAL_DEPENDENCIES = {
    "pdf_test_fixture": ("reportlab", "reportlab"),
    "pptx_import": ("python-pptx", "pptx"),
    "xlsx_import": ("openpyxl", "openpyxl"),
}
CONTROLLED_PILOT_ENV = "NEXUS_CONTROLLED_PILOT_MODE"
CONTROLLED_PILOT_TITLE = "PILOTE LOCAL — VALIDATION HUMAINE OBLIGATOIRE"
CONTROLLED_PILOT_NOTICE = (
    "Cette analyse est une aide à la préparation syndicale. "
    "Elle doit être vérifiée avant toute utilisation auprès d’un salarié, "
    "de l’employeur ou d’une instance."
)


class NexusHTTPServer(ThreadingHTTPServer):
    """Local server that refuses a second process on the same Windows port."""

    allow_reuse_address = False

    def server_bind(self) -> None:
        if hasattr(socket, "SO_EXCLUSIVEADDRUSE"):
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()
MAX_FOLLOW_UP_ANSWERS = 12
MAX_FOLLOW_UP_QUESTION_LENGTH = 500
MAX_FOLLOW_UP_ANSWER_LENGTH = 3000
MAX_QUESTION_DOCUMENTS = 3
MAX_QUESTION_DOCUMENT_BYTES = 5 * 1024 * 1024
MAX_QUESTION_DOCUMENT_BASE64_CHARS = 7 * 1024 * 1024
MAX_QUESTION_DOCUMENT_CHARS = 20000
MAX_ALL_QUESTION_DOCUMENT_CHARS = 40000
QUESTION_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".jpg", ".jpeg"}
LOCAL_CASES_ROOT = Path(
    os.getenv(
        "NEXUS_LOCAL_CASES_ROOT",
        str(ROOT / "local-index" / "nexus-local-cases"),
    )
)
LOCAL_CASE_STORE = LocalCaseStore(LOCAL_CASES_ROOT)

PORTAL_CONTEXT_LABELS = {
    "startDate": "Date de début",
    "eventFrequency": "Fréquence de la situation",
    "stillEmployed": "Le salarié est encore en poste",
    "procedureOngoing": "Une procédure est en cours",
    "urgentSituation": "La situation est signalée comme urgente",
    "knownDeadline": "Échéance connue",
    "employeeCount": "Nombre approximatif de salariés",
    "services": "Services concernés",
    "decisionAnnounced": "La décision a déjà été annoncée",
    "implementationStarted": "La mise en œuvre a commencé",
    "meetingPlanned": "Une réunion est prévue",
    "meetingDate": "Date de réunion",
    "payrollMonth": "Mois de paie concerné",
    "periodDetails": "Période précisée",
    "recurringGap": "L’écart est récurrent",
    "alreadyReported": "L’anomalie a déjà été signalée",
}


def _env_enabled(name: str) -> bool:
    return os.getenv(name, "").strip().lower() in {"1", "true", "yes", "on"}


def controlled_pilot_payload() -> dict[str, object]:
    """Return the public, non-sensitive local pilot notice."""

    enabled = _env_enabled(CONTROLLED_PILOT_ENV)
    return {
        "enabled": enabled,
        "title": CONTROLLED_PILOT_TITLE if enabled else "",
        "notice": CONTROLLED_PILOT_NOTICE if enabled else "",
    }


def _bounded_text(value: object, limit: int) -> str:
    return " ".join(str(value or "").split())[:limit].strip()


def _safe_upload_name(value: object) -> str:
    name = Path(str(value or "document").replace("\\", "/")).name
    return re.sub(r"[^0-9A-Za-zÀ-ÖØ-öø-ÿ._() -]", "_", name)[:160] or "document"


def _clean_document_text(value: object) -> str:
    return "\n".join(
        line.strip() for line in str(value or "").splitlines() if line.strip()
    )


def _docx_list_metadata(document: object, paragraph: object) -> tuple[str, int, int] | None:
    """Return list format, list id and level for a Word paragraph."""

    try:
        from docx.oxml.ns import qn

        paragraph_properties = paragraph._p.pPr
        numbering = paragraph_properties.numPr if paragraph_properties is not None else None
        if numbering is None and paragraph.style is not None:
            style_properties = paragraph.style.element.pPr
            numbering = style_properties.numPr if style_properties is not None else None
        if numbering is None or numbering.numId is None:
            return None
        list_id = int(numbering.numId.val)
        level = int(numbering.ilvl.val) if numbering.ilvl is not None else 0
        numbering_root = document.part.numbering_part.element
        abstract_id: int | None = None
        for instance in numbering_root.findall(qn("w:num")):
            if int(instance.get(qn("w:numId"))) == list_id:
                abstract = instance.find(qn("w:abstractNumId"))
                if abstract is not None:
                    abstract_id = int(abstract.get(qn("w:val")))
                break
        if abstract_id is None:
            return ("decimal", list_id, level)
        for abstract in numbering_root.findall(qn("w:abstractNum")):
            if int(abstract.get(qn("w:abstractNumId"))) != abstract_id:
                continue
            for level_node in abstract.findall(qn("w:lvl")):
                if int(level_node.get(qn("w:ilvl"))) != level:
                    continue
                number_format = level_node.find(qn("w:numFmt"))
                value = number_format.get(qn("w:val")) if number_format is not None else "decimal"
                return (str(value), list_id, level)
        return ("decimal", list_id, level)
    except (AttributeError, KeyError, TypeError, ValueError):
        return None


def _docx_passages(document: object) -> list[dict[str, str]]:
    """Extract Word paragraphs while preserving automatic agenda numbering."""

    passages: list[dict[str, str]] = []
    counters: dict[tuple[int, int], int] = {}
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = _clean_document_text(paragraph.text)
        if not text:
            continue
        metadata = _docx_list_metadata(document, paragraph)
        if metadata is not None:
            number_format, list_id, level = metadata
            if number_format == "bullet" or level > 0:
                text = f"- {text}"
            else:
                key = (list_id, level)
                counters[key] = counters.get(key, 0) + 1
                text = f"{counters[key]}. {text}"
        passages.append({"reference": f"Paragraphe {index}", "text": text})
    return passages


def _extract_uploaded_document(raw_document: object) -> dict[str, Any]:
    if not isinstance(raw_document, dict):
        raise ValueError("Pièce jointe invalide.")
    name = _safe_upload_name(raw_document.get("name"))
    extension = Path(name).suffix.casefold()
    if extension not in QUESTION_DOCUMENT_EXTENSIONS:
        raise ValueError(
            f"Format non accepté pour {name}. Utilisez PDF, DOCX, TXT, MD, JPG ou JPEG."
        )
    encoded = str(raw_document.get("content_base64") or "")
    if len(encoded) > MAX_QUESTION_DOCUMENT_BASE64_CHARS:
        raise ValueError(f"Le document {name} dépasse 5 Mo.")
    try:
        content = base64.b64decode(encoded, validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"Le document {name} est illisible.") from exc
    if not content or len(content) > MAX_QUESTION_DOCUMENT_BYTES:
        raise ValueError(f"Le document {name} est vide ou dépasse 5 Mo.")
    passages: list[dict[str, str]] = []
    page_count: int | None = None
    paragraph_count: int | None = None
    if extension in {".txt", ".md"}:
        text = content.decode("utf-8", errors="replace")
        raw_passages = [part for part in re.split(r"\n\s*\n", text) if part.strip()]
        passages = [
            {"reference": f"Paragraphe {index}", "text": _clean_document_text(part)}
            for index, part in enumerate(raw_passages, start=1)
            if _clean_document_text(part)
        ]
        paragraph_count = len(passages)
    elif extension in {".jpg", ".jpeg"}:
        try:
            passages = extract_text_from_jpeg(content)
        except LocalOCRUnavailable as exc:
            raise ValueError(f"Le document {name} ne peut pas être lu : {exc}") from exc
        except ValueError as exc:
            raise ValueError(f"Le document {name} ne peut pas être lu : {exc}") from exc
        page_count = 1
        paragraph_count = len(passages)
    elif extension == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            page_count = len(reader.pages)
            for index, page in enumerate(reader.pages, start=1):
                page_text = _clean_document_text(page.extract_text())
                if page_text:
                    passages.append({"reference": f"Page {index}", "text": page_text})
        except Exception as exc:
            raise ValueError(f"Le texte du PDF {name} ne peut pas être lu.") from exc
    else:
        try:
            from docx import Document

            document = Document(io.BytesIO(content))
            passages = _docx_passages(document)
            paragraph_count = len(passages)
            for table_index, table in enumerate(document.tables, start=1):
                table_text = _clean_document_text(
                    "\n".join(cell.text for row in table.rows for cell in row.cells)
                )
                if table_text:
                    passages.append(
                        {"reference": f"Tableau {table_index}", "text": table_text}
                    )
        except Exception as exc:
            raise ValueError(f"Le texte du document Word {name} ne peut pas être lu.") from exc
    normalized = "\n\n".join(item["text"] for item in passages if item["text"])
    if not normalized:
        raise ValueError(
            f"Aucun texte exploitable n’a été trouvé dans {name}. "
            "S’il s’agit d’un document scanné, collez le passage utile dans la question."
        )
    bounded_text = normalized[:MAX_QUESTION_DOCUMENT_CHARS]
    return {
        "name": name,
        "text": bounded_text,
        "extension": extension,
        "size_bytes": len(content),
        "page_count": page_count,
        "paragraph_count": paragraph_count,
        "truncated": len(normalized) > len(bounded_text),
        "passages": passages,
    }


def extract_uploaded_documents(raw_documents: object) -> list[dict[str, Any]]:
    if raw_documents in (None, []):
        return []
    if not isinstance(raw_documents, list) or len(raw_documents) > MAX_QUESTION_DOCUMENTS:
        raise ValueError("Trois pièces jointes maximum sont acceptées.")
    documents = [_extract_uploaded_document(item) for item in raw_documents]
    remaining = MAX_ALL_QUESTION_DOCUMENT_CHARS
    bounded: list[dict[str, Any]] = []
    for document in documents:
        text = document["text"][:remaining]
        if text:
            bounded.append(
                {
                    **document,
                    "text": text,
                    "truncated": bool(document["truncated"] or len(document["text"]) > len(text)),
                }
            )
            remaining -= len(text)
        if remaining <= 0:
            break
    return bounded


UPLOAD_STOPWORDS = {
    "avec", "dans", "document", "fichier", "pour", "mais", "plus", "peut",
    "cette", "comme", "elle", "entre", "sera", "sont", "avoir", "faire",
    "salarié", "salariée", "employeur", "question", "joint", "jointe",
}


def _search_tokens(value: object) -> set[str]:
    folded = unicodedata.normalize("NFKD", str(value or ""))
    folded = "".join(char for char in folded if not unicodedata.combining(char))
    return {
        token
        for token in re.findall(r"[a-z0-9]{4,}", folded.casefold())
        if token not in UPLOAD_STOPWORDS
    }


def _best_uploaded_passage(
    document: dict[str, Any],
    relevance_text: str,
) -> tuple[str, str]:
    wanted = _search_tokens(relevance_text)
    ranked: list[tuple[int, int, str, str]] = []
    for index, passage in enumerate(document.get("passages") or []):
        if not isinstance(passage, dict):
            continue
        text = _clean_document_text(passage.get("text"))
        if not text:
            continue
        score = len(wanted & _search_tokens(text))
        ranked.append((score, -index, str(passage.get("reference") or ""), text))
    if not ranked:
        return "", ""
    _, _, reference, text = max(ranked, key=lambda item: (item[0], item[1]))
    if len(text) <= 900:
        return reference, text
    sentences = [part.strip() for part in re.split(r"(?<=[.!?])\s+", text) if part.strip()]
    if not sentences:
        return reference, text[:900].rstrip() + "…"
    best_index = max(
        range(len(sentences)),
        key=lambda index: len(wanted & _search_tokens(sentences[index])),
    )
    selected = sentences[best_index]
    for distance in range(1, len(sentences)):
        for candidate_index in (best_index - distance, best_index + distance):
            if candidate_index < 0 or candidate_index >= len(sentences):
                continue
            candidate = (
                f"{sentences[candidate_index]} {selected}"
                if candidate_index < best_index
                else f"{selected} {sentences[candidate_index]}"
            )
            if len(candidate) <= 900:
                selected = candidate
        if len(selected) >= 650:
            break
    return reference, selected[:900].rstrip()


def add_uploaded_documents_to_public_result(
    result: dict[str, Any],
    documents: list[dict[str, Any]],
    user_question: str,
) -> dict[str, Any]:
    """Expose exact, bounded user-document excerpts as factual evidence.

    A user document is never labelled as a legal norm. Full extracted text and
    binary content remain request-local and are not returned to the browser.
    """

    if not documents:
        return result
    answer = result.get("answer") if isinstance(result.get("answer"), dict) else {}
    core = answer.get("case_factual_core") if isinstance(answer.get("case_factual_core"), dict) else {}
    relevance_text = " ".join(
        str(value or "")
        for value in (
            user_question,
            core.get("event_category"),
            core.get("primary_grievance_or_decision"),
            *(core.get("facts_certain") or []),
        )
    )
    projections: list[dict[str, Any]] = []
    summaries: list[dict[str, Any]] = []
    for document in documents:
        reference, excerpt = _best_uploaded_passage(document, relevance_text)
        if not excerpt:
            continue
        title = _safe_upload_name(document.get("name"))
        truncated = bool(document.get("truncated"))
        projections.append(
            {
                "provider": "Document fourni par l’utilisateur",
                "title": title,
                "nature": "USER_DOCUMENT",
                "availability_status": "PROVIDED_AND_READ",
                "reference": reference or "Passage extrait",
                "excerpt": excerpt,
                "link_to_facts": "Ce passage a été utilisé comme élément factuel du dossier.",
                "practical_scope": (
                    "Pièce factuelle fournie pour cette analyse ; elle ne constitue pas, "
                    "à elle seule, une règle juridique."
                ),
                "reserve": (
                    "Le document a été partiellement lu en raison de sa longueur."
                    if truncated
                    else "Vérifier l’authenticité, la date et la version du document."
                ),
            }
        )
        summaries.append(
            {
                "title": title,
                "format": str(document.get("extension") or "").removeprefix(".").upper(),
                "page_count": document.get("page_count"),
                "paragraph_count": document.get("paragraph_count"),
                "truncated": truncated,
                "status": "LU_POUR_CETTE_ANALYSE",
            }
        )
    if not projections:
        return result
    public_summary = result.get("public_summary")
    if not isinstance(public_summary, dict):
        public_summary = {}
        result["public_summary"] = public_summary
    existing_extractions = public_summary.get("source_extractions")
    public_summary["source_extractions"] = [
        *projections,
        *(existing_extractions if isinstance(existing_extractions, list) else []),
    ]
    public_summary["uploaded_documents"] = summaries
    detailed = result.get("detailed_analysis")
    if isinstance(detailed, dict):
        detailed["uploaded_documents"] = summaries
    return result


def _credential_pair_configured(*pairs: tuple[str, str]) -> bool:
    return any(bool(os.getenv(client_id) and os.getenv(client_secret)) for client_id, client_secret in pairs)


def runtime_capability_status() -> dict[str, Any]:
    """Return only non-sensitive configuration availability for the local UI."""

    cse_root = str(os.getenv("NEXUS_CSE_MEMORY_PROCESSED_ROOT") or "").strip()
    protection_root = str(os.getenv("NEXUS_PROTECTION_SOCIALE_PROCESSED_ROOT") or "").strip()
    return {
        "runtime": {
            "core": _env_enabled("NEXUS_CORE_RUNTIME_ENABLED"),
            "connectors": _env_enabled("NEXUS_CONNECTOR_RUNTIME_ENABLED"),
            "cse_memory": _env_enabled("NEXUS_CSE_MEMORY_RUNTIME_ENABLED"),
            "official_connectors": _env_enabled("NEXUS_OFFICIAL_CONNECTORS_RUNTIME_ENABLED"),
            "syndical_reasoning": _env_enabled("NEXUS_SYNDICAL_REASONING_RUNTIME_ENABLED"),
            "source_execution": _env_enabled("NEXUS_SOURCE_EXECUTION_COORDINATOR_ENABLED"),
            "source_execution_network": _env_enabled("NEXUS_SOURCE_EXECUTION_NETWORK_ENABLED"),
            "retrieval_to_response": _env_enabled("NEXUS_RETRIEVAL_TO_FINAL_RESPONSE_ENABLED"),
        },
        "network": {
            "official_knowledge": _env_enabled("OFFICIAL_KNOWLEDGE_NETWORK_ENABLED"),
        },
        "credentials": {
            "legifrance": _credential_pair_configured(
                ("CFDT_NEXUS_LEGIFRANCE_CLIENT_ID", "CFDT_NEXUS_LEGIFRANCE_CLIENT_SECRET"),
                ("LEGIFRANCE_CLIENT_ID", "LEGIFRANCE_CLIENT_SECRET"),
                ("PISTE_CLIENT_ID", "PISTE_CLIENT_SECRET"),
            ),
            "judilibre": _credential_pair_configured(
                ("CFDT_NEXUS_JUDILIBRE_CLIENT_ID", "CFDT_NEXUS_JUDILIBRE_CLIENT_SECRET"),
                ("JUDILIBRE_CLIENT_ID", "JUDILIBRE_CLIENT_SECRET"),
                ("CFDT_NEXUS_LEGIFRANCE_CLIENT_ID", "CFDT_NEXUS_LEGIFRANCE_CLIENT_SECRET"),
                ("PISTE_CLIENT_ID", "PISTE_CLIENT_SECRET"),
            ),
        },
        "local_corpora": {
            "cse_memory": bool(cse_root and Path(cse_root).is_dir()),
            "protection_sociale": bool(protection_root and Path(protection_root).is_dir()),
        },
    }


def _follow_up_rows(portal_context: dict[str, Any]) -> list[tuple[str, str]]:
    raw_rows = portal_context.get("follow_up_answers")
    if not isinstance(raw_rows, list):
        return []
    rows: list[tuple[str, str]] = []
    for raw_row in raw_rows[:MAX_FOLLOW_UP_ANSWERS]:
        if not isinstance(raw_row, dict):
            continue
        question = _bounded_text(
            raw_row.get("question"), MAX_FOLLOW_UP_QUESTION_LENGTH
        )
        answer = _bounded_text(raw_row.get("answer"), MAX_FOLLOW_UP_ANSWER_LENGTH)
        if question and answer:
            rows.append((question, answer))
    return rows


def build_enriched_analysis_query(
    raw_query: str,
    portal_context: dict[str, Any],
) -> str:
    """Project portal answers into the factual-core structured input contract.

    The projection is request-local: no answer is persisted or shared with another
    analysis. Employee statements remain allegations until corroborated.
    """

    user_question = _bounded_text(
        portal_context.get("user_question"), MAX_FOLLOW_UP_ANSWER_LENGTH
    )
    interview_rows = portal_context.get("employee_interview")
    follow_up_rows = _follow_up_rows(portal_context)
    uploaded_documents = portal_context.get("uploaded_documents")
    if (
        not user_question
        and not isinstance(interview_rows, list)
        and not follow_up_rows
        and not isinstance(uploaded_documents, list)
    ):
        return raw_query

    facts = [user_question] if user_question else [_bounded_text(raw_query, 6000)]
    if isinstance(uploaded_documents, list):
        for document in uploaded_documents[:MAX_QUESTION_DOCUMENTS]:
            if not isinstance(document, dict):
                continue
            name = _safe_upload_name(document.get("name"))
            text = str(document.get("text") or "").strip()
            if text:
                factual_text = re.sub(r"\s*:\s*", " — ", text)
                facts.append(
                    f"Contenu factuel du document {name} — {factual_text}"
                )
    context_rows: list[str] = []
    raw_context = portal_context.get("facts")
    if isinstance(raw_context, dict):
        for key, value in raw_context.items():
            if value in (None, "", False):
                continue
            label = PORTAL_CONTEXT_LABELS.get(str(key))
            if not label:
                continue
            rendered = "oui" if value is True else _bounded_text(value, 500)
            context_rows.append(f"{label} : {rendered}")

    allegations: list[str] = []
    if isinstance(interview_rows, list):
        for raw_row in interview_rows[:MAX_FOLLOW_UP_ANSWERS]:
            if not isinstance(raw_row, dict):
                continue
            question = _bounded_text(
                raw_row.get("question"), MAX_FOLLOW_UP_QUESTION_LENGTH
            )
            answer = _bounded_text(
                raw_row.get("answer"), MAX_FOLLOW_UP_ANSWER_LENGTH
            )
            if question and answer:
                allegations.append(
                    f"En réponse à « {question} », le salarié indique : {answer}"
                )
    allegations.extend(
        f"En réponse à « {question} », le salarié indique : {answer}"
        for question, answer in follow_up_rows
    )

    documents = portal_context.get("available_documents")
    if isinstance(documents, list):
        names = [
            _bounded_text(name, 180)
            for name in documents[:20]
            if _bounded_text(name, 180)
        ]
        if names:
            context_rows.append("Documents déclarés disponibles : " + ", ".join(names))

    sections = ["Faits fournis:", *(f"- {item}" for item in facts if item)]
    if allegations:
        sections.extend(
            ["Faits allégués:", *(f"- {item}" for item in allegations)]
        )
    if context_rows:
        sections.extend(["Contexte:", *(f"- {item}" for item in context_rows)])
    return "\n".join(sections)


def optional_dependency_status() -> dict[str, dict[str, object]]:
    """Report optional capabilities without importing or installing packages."""

    return {
        capability: {
            "package": package,
            "available": importlib.util.find_spec(module) is not None,
        }
        for capability, (package, module) in OPTIONAL_DEPENDENCIES.items()
    }


def router_environment() -> dict[str, str]:
    env = os.environ.copy()
    existing_pythonpath = env.get("PYTHONPATH", "")
    env["PYTHONPATH"] = str(ROOT) + (
        os.pathsep + existing_pythonpath if existing_pythonpath else ""
    )
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    return env


def run_router(
    query: str,
    source_limit: int = 6,
    employee_path: str | None = None,
) -> dict[str, Any]:
    command = [
        sys.executable,
        "-B",
        str(ROUTER_SCRIPT),
        "ask",
        "--query",
        query,
        "--source-limit",
        str(source_limit),
        "--format",
        "json",
    ]
    if employee_path:
        command.extend(["--employee-path", employee_path])
    completed = subprocess.run(
        command,
        cwd=str(ROOT),
        env=router_environment(),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        timeout=180,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or "Assistant DS Router indisponible.")
    return json.loads(completed.stdout)


def analyze_question(
    query: str,
    source_limit: int = 6,
    employee_path: str | None = None,
) -> dict[str, Any]:
    """Build the complete internal payload used by Runtime integrations."""

    cleaned = (query or "").strip()
    if not cleaned:
        raise ValueError("Question vide.")
    answer = (
        run_router(cleaned, source_limit, employee_path)
        if employee_path
        else run_router(cleaned, source_limit)
    )
    connector_config = RuntimeConnectorConfig.from_env()
    connector_mapping = RuntimeConnectorPayloadMapper(connector_config).map(answer)
    official_config = RuntimeOfficialConnectorsConfig.from_env()
    official_connectors = RuntimeOfficialConnectorsIntegration(official_config).integrate(answer)
    if official_connectors.questions or official_connectors.source_qualifications:
        answer = dict(answer)
        answer["questions_to_ask"] = list(
            dict.fromkeys(
                (
                    *(
                        answer.get("questions_to_ask", ())
                        if isinstance(answer.get("questions_to_ask"), list)
                        else ()
                    ),
                    *official_connectors.questions,
                )
            )
        )
        answer["official_source_qualifications"] = [
            qualification.to_dict()
            for qualification in official_connectors.source_qualifications
        ]
        answer["source_extraction"] = merge_metadata_source_qualifications(
            answer.get("source_extraction"),
            answer["official_source_qualifications"],
        )
    expert_payload = orchestrator.orchestrate(answer)
    payload = {
        "ok": True,
        "answer": answer,
        **expert_payload,
    }
    payload["official_connectors_runtime"] = official_connectors.to_dict()
    integration = RuntimeCoreIntegration(RuntimeIntegrationConfig.from_env()).integrate(
        RuntimeCoreIntegrationInput(
            answer=answer,
            legal_payload=payload.get("expert_juriste"),
            payroll_payload=payload.get("expert_paie"),
            historical_orchestration=payload.get("orchestration") or {},
            connector_inputs=connector_mapping.inputs + official_connectors.inputs,
            connector_runtime_enabled=connector_config.enabled or official_config.enabled,
            connector_mapping_fallback_code=connector_mapping.fallback_code,
        )
    )
    payload["runtime_integration"] = integration.to_dict()
    historical_report = report_generator.build_report(payload)
    core_report = RuntimeCoreReportMapper().map(historical_report, integration)
    cse_config = RuntimeCSEMemoryConfig.from_env(
        default_root=ROOT / "CCSEMEMORYENGINE" / "PROCESSED" / "LOT_1D"
    )
    try:
        cse_integration = RuntimeCSEMemoryIntegration(cse_config).integrate(answer)
    except Exception:
        cse_integration = RuntimeCSEMemoryResult(
            RuntimeCSEMemoryMode.FALLBACK,
            RuntimeCSEMemoryDiagnostics(
                cse_config.enabled,
                called=True,
                fallback_triggered=True,
                fallback_code="CSE_MEMORY_RUNTIME_FAILED",
            ),
        )
    payload["cse_memory_runtime"] = cse_integration.to_dict()
    cse_report = RuntimeCSEMemoryReportMapper().map(core_report, cse_integration)
    retirement_integration = RuntimeRetirementIntegration(
        RuntimeRetirementConfig.from_env()
    ).integrate(answer)
    payload["retirement_runtime"] = retirement_integration.to_dict()
    retirement_report = RuntimeRetirementReportMapper().map(
        cse_report, retirement_integration
    )
    protection_config = RuntimeProtectionSocialeConfig.from_env(
        default_root=ROOT / "PROTECTION_SOCIALE_ENGINE" / "PROCESSED" / "LOT_1D"
    )
    protection_integration = RuntimeProtectionSocialeIntegration(
        protection_config
    ).integrate(answer)
    payload["protection_sociale_runtime"] = protection_integration.to_dict()
    protection_report = RuntimeProtectionSocialeReportMapper().map(
        retirement_report, protection_integration
    )
    syndical_integration = RuntimeSyndicalReasoningIntegration(
        RuntimeSyndicalReasoningConfig.from_env()
    ).integrate(answer)
    payload["syndical_reasoning_runtime"] = syndical_integration.to_dict()
    syndical_report = RuntimeSyndicalReasoningReportMapper().map(
        protection_report, syndical_integration
    )
    final_config = RuntimeFinalAssistantConfig.from_env()
    if final_config.enabled:
        from NEXUS_RUNTIME_INTEGRATION.final_assistant_runtime import (
            RuntimeFinalAssistantIntegration,
        )

        final_integration = RuntimeFinalAssistantIntegration(final_config).integrate(
            answer,
            syndical_report,
            existing_results={
                "syndical_reasoning": syndical_integration.to_dict(),
                "cse_memory": cse_integration.to_dict(),
            },
        )
        payload["final_assistant_runtime"] = final_integration.to_dict()
        payload["analysis_report"] = final_integration.report
    else:
        payload["final_assistant_runtime"] = {
            "mode": "DISABLED",
            "diagnostics": {
                "enabled": False,
                "called": False,
                "runtime_ms": 0,
                "engines_used": [],
                "fallback_code": None,
            },
            "assistant": None,
        }
        payload["analysis_report"] = syndical_report
    return payload


def analyze_public_question(
    query: str,
    source_limit: int = 6,
    employee_path: str | None = None,
) -> dict[str, Any]:
    """Return the user-facing payload without internal identifiers or paths."""

    result = (
        analyze_question(query, source_limit, employee_path)
        if employee_path
        else analyze_question(query, source_limit)
    )
    return sanitize_public_payload(result)


class NexusHandler(SimpleHTTPRequestHandler):
    server_version = "NexusLocalInterface"

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, directory=str(APP_DIR), **kwargs)

    def log_message(self, format: str, *args: Any) -> None:  # noqa: A003
        try:
            sys.stderr.write("[nexus-local] " + format % args + "\n")
            sys.stderr.flush()
        except (OSError, ValueError):
            # A desktop launcher may close its console after detaching Nexus.
            # Request handling must never depend on that diagnostic stream.
            pass

    def end_headers(self) -> None:
        self.send_header("Cache-Control", "no-store, max-age=0")
        self.send_header("Pragma", "no-cache")
        super().end_headers()

    def send_json(self, status: HTTPStatus, payload: dict[str, Any]) -> None:
        outgoing = dict(payload)
        pilot = controlled_pilot_payload()
        if pilot["enabled"]:
            outgoing["controlled_pilot"] = pilot
        safe_payload = sanitize_http_public_payload(outgoing)
        data = json.dumps(safe_payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def send_internal_error(self, exc: Exception) -> None:
        self.log_error("Internal server error (%s)", type(exc).__name__)
        self.send_json(HTTPStatus.INTERNAL_SERVER_ERROR, {"ok": False, "error": INTERNAL_ERROR_MESSAGE})

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path in {"/", ""}:
            self.path = "/index.html"
        if parsed.path == "/health":
            self.send_json(
                HTTPStatus.OK,
                {
                    "ok": True,
                    "service": "nexus-local-interface",
                    "version": get_nexus_version(),
                    "mode": "local",
                    "persistent_case_storage": False,
                    "voluntary_local_case_storage": True,
                    "controlled_pilot": controlled_pilot_payload(),
                    "optional_dependencies": optional_dependency_status(),
                    "runtime_capabilities": runtime_capability_status(),
                },
            )
            return
        if parsed.path == "/api/employee-case/scenarios":
            self.send_json(HTTPStatus.OK, {"ok": True, "scenarios": public_scenarios(), "synthetic_only": True})
            return
        if parsed.path == "/api/employee-case/demo":
            try:
                scenario = (parse_qs(parsed.query).get("scenario") or [""])[0]
                if not scenario:
                    raise ValueError("Scenario synthetique manquant.")
                self.send_json(HTTPStatus.OK, build_demo_payload(scenario))
            except KeyError as exc:
                self.send_json(HTTPStatus.NOT_FOUND, {"ok": False, "error": str(exc.args[0])})
            except PermissionError as exc:
                self.send_json(HTTPStatus.FORBIDDEN, {"ok": False, "error": str(exc)})
            except ValueError as exc:
                self.send_json(HTTPStatus.BAD_REQUEST, {"ok": False, "error": str(exc)})
            except Exception as exc:  # pragma: no cover - defensive local server boundary.
                self.send_internal_error(exc)
            return
        if parsed.path == "/api/historical-cases":
            try:
                query = (parse_qs(parsed.query).get("query") or [""])[0]
                category = (parse_qs(parsed.query).get("category") or ["all"])[0]
                self.send_json(
                    HTTPStatus.OK,
                    {
                        "ok": True,
                        **list_historical_cases(query=query, category=category),
                    },
                )
            except ValueError as exc:
                self.send_json(
                    HTTPStatus.BAD_REQUEST,
                    {"ok": False, "error": str(exc)},
                )
            except Exception as exc:  # pragma: no cover - defensive local boundary.
                self.send_internal_error(exc)
            return
        if parsed.path == "/api/rgpd/corpus-scan":
            try:
                self.send_json(
                    HTTPStatus.OK,
                    {"ok": True, "rgpd_analysis": build_rgpd_corpus_scan()},
                )
            except Exception as exc:  # pragma: no cover - defensive local server boundary.
                self.send_internal_error(exc)
            return
        if parsed.path == "/api/local-cases":
            try:
                self.send_json(
                    HTTPStatus.OK,
                    {"ok": True, "cases": LOCAL_CASE_STORE.list_cases()},
                )
            except Exception as exc:  # pragma: no cover - defensive local boundary.
                self.send_internal_error(exc)
            return
        if parsed.path.startswith("/api/local-cases/"):
            try:
                case_ref = parsed.path.removeprefix("/api/local-cases/")
                self.send_json(
                    HTTPStatus.OK,
                    {"ok": True, "case": LOCAL_CASE_STORE.get_case(case_ref)},
                )
            except KeyError as exc:
                self.send_json(
                    HTTPStatus.NOT_FOUND,
                    {"ok": False, "error": str(exc.args[0])},
                )
            except ValueError as exc:
                self.send_json(
                    HTTPStatus.BAD_REQUEST,
                    {"ok": False, "error": str(exc)},
                )
            except Exception as exc:  # pragma: no cover - defensive local boundary.
                self.send_internal_error(exc)
            return
        if (
            parsed.path.startswith("/api/historical-cases/")
            and parsed.path.endswith("/refresh-sources")
        ):
            try:
                case_id = parsed.path.removeprefix(
                    "/api/historical-cases/"
                ).removesuffix("/refresh-sources")
                self.send_json(
                    HTTPStatus.OK,
                    {
                        "ok": True,
                        "refresh": refresh_historical_case_sources(case_id),
                    },
                )
            except KeyError as exc:
                self.send_json(
                    HTTPStatus.NOT_FOUND,
                    {"ok": False, "error": str(exc.args[0])},
                )
            except (SystemExit, ValueError) as exc:
                self.send_json(
                    HTTPStatus.BAD_REQUEST,
                    {"ok": False, "error": str(exc)},
                )
            except Exception as exc:  # pragma: no cover - defensive local boundary.
                self.send_internal_error(exc)
            return
        if parsed.path.startswith("/api/historical-cases/"):
            try:
                case_id = parsed.path.removeprefix("/api/historical-cases/")
                self.send_json(
                    HTTPStatus.OK,
                    {"ok": True, "case": get_historical_case(case_id)},
                )
            except KeyError as exc:
                self.send_json(
                    HTTPStatus.NOT_FOUND,
                    {"ok": False, "error": str(exc.args[0])},
                )
            except ValueError as exc:
                self.send_json(
                    HTTPStatus.BAD_REQUEST,
                    {"ok": False, "error": str(exc)},
                )
            except Exception as exc:  # pragma: no cover - defensive local boundary.
                self.send_internal_error(exc)
            return
        super().do_GET()

    def do_POST(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path not in {
            "/api/analyze",
            "/api/local-cases",
            "/api/cse/agenda-preview",
            "/api/rgpd/analyze",
        }:
            self.send_json(HTTPStatus.NOT_FOUND, {"ok": False, "error": "Endpoint inconnu."})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            body = self.rfile.read(length).decode("utf-8")
            payload = json.loads(body or "{}")
            if parsed.path == "/api/local-cases":
                self.send_json(
                    HTTPStatus.OK,
                    {"ok": True, "case": LOCAL_CASE_STORE.save_case(payload)},
                )
                return
            if parsed.path == "/api/rgpd/analyze":
                uploaded_documents = extract_uploaded_documents(payload.get("attachments"))
                if not uploaded_documents:
                    raise ValueError("Aucun document a analyser.")
                self.send_json(
                    HTTPStatus.OK,
                    {"ok": True, "rgpd_analysis": build_rgpd_analysis(uploaded_documents)},
                )
                return
            if parsed.path == "/api/cse/agenda-preview":
                uploaded_documents = extract_uploaded_documents(payload.get("attachments"))
                previous_minutes_documents = extract_uploaded_documents(
                    payload.get("cse_previous_pv_attachments")
                )
                self.send_json(
                    HTTPStatus.OK,
                    {
                        "ok": True,
                        "cse_agenda_analysis": build_cse_agenda_analysis(
                            uploaded_documents,
                            additional_minutes_documents=previous_minutes_documents,
                        ),
                    },
                )
                return
            source_limit = int(payload.get("source_limit") or 6)
            portal_context = (
                payload.get("portal_context")
                if isinstance(payload.get("portal_context"), dict)
                else {}
            )
            uploaded_documents = extract_uploaded_documents(payload.get("attachments"))
            previous_minutes_documents = (
                extract_uploaded_documents(payload.get("cse_previous_pv_attachments"))
                if portal_context.get("workspace") == "cse"
                else []
            )
            if uploaded_documents:
                portal_context = dict(portal_context)
                portal_context["uploaded_documents"] = uploaded_documents
            employee_path = payload.get("employee_path") or portal_context.get(
                "employee_path"
            )
            analysis_query = build_enriched_analysis_query(
                str(payload.get("query") or ""),
                portal_context,
            )
            result = analyze_public_question(
                analysis_query,
                source_limit,
                str(employee_path) if employee_path else None,
            )
            result = add_uploaded_documents_to_public_result(
                result,
                uploaded_documents,
                str(portal_context.get("user_question") or payload.get("query") or ""),
            )
            if portal_context.get("workspace") == "cse":
                selected_positions = portal_context.get("cse_selected_agenda_positions")
                result["cse_agenda_analysis"] = build_cse_agenda_analysis(
                    uploaded_documents,
                    selected_positions=(
                        selected_positions
                        if isinstance(selected_positions, list)
                        else None
                    ),
                    additional_minutes_documents=previous_minutes_documents,
                )
            self.send_json(HTTPStatus.OK, result)
        except ValueError as exc:
            self.send_json(HTTPStatus.BAD_REQUEST, {"ok": False, "error": str(exc)})
        except Exception as exc:  # pragma: no cover - defensive local server boundary.
            self.send_internal_error(exc)

    def do_DELETE(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if not parsed.path.startswith("/api/local-cases/"):
            self.send_json(
                HTTPStatus.NOT_FOUND,
                {"ok": False, "error": "Endpoint inconnu."},
            )
            return
        try:
            case_ref = parsed.path.removeprefix("/api/local-cases/")
            LOCAL_CASE_STORE.delete_case(case_ref)
            self.send_json(HTTPStatus.OK, {"ok": True})
        except KeyError as exc:
            self.send_json(
                HTTPStatus.NOT_FOUND,
                {"ok": False, "error": str(exc.args[0])},
            )
        except ValueError as exc:
            self.send_json(
                HTTPStatus.BAD_REQUEST,
                {"ok": False, "error": str(exc)},
            )
        except Exception as exc:  # pragma: no cover - defensive local boundary.
            self.send_internal_error(exc)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="CFDT Nexus - interface locale")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--open", action="store_true", help="ouvrir le navigateur local")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    server = NexusHTTPServer((args.host, args.port), NexusHandler)
    url = f"http://{args.host}:{args.port}/"
    print(f"Nexus interface locale: {url}")
    print("Aucun acces internet requis. Arreter avec Ctrl+C.")
    if args.open:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nArret de Nexus local.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
