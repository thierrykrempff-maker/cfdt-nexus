from __future__ import annotations

import base64
import importlib.util
import io
from pathlib import Path
import sys

import pytest
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SERVER_PATH = ROOT / "apps" / "nexus-local-interface" / "server.py"
sys.path.insert(0, str(SERVER_PATH.parent))
SPEC = importlib.util.spec_from_file_location("nexus_local_upload_server", SERVER_PATH)
assert SPEC and SPEC.loader
SERVER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(SERVER)


def attachment(name: str, content: bytes, mime_type: str = "text/plain") -> dict[str, str]:
    return {
        "name": name,
        "mime_type": mime_type,
        "content_base64": base64.b64encode(content).decode("ascii"),
    }


def test_text_document_is_extracted_in_memory_and_added_to_analysis() -> None:
    documents = SERVER.extract_uploaded_documents(
        [attachment("avenant.txt", "Passage en poste avec week-ends.".encode("utf-8"))]
    )
    query = SERVER.build_enriched_analysis_query(
        "Le salarié peut-il refuser ?",
        {
            "user_question": "Le salarié peut-il refuser ?",
            "uploaded_documents": documents,
        },
    )
    assert documents[0]["name"] == "avenant.txt"
    assert documents[0]["text"] == "Passage en poste avec week-ends."
    assert documents[0]["passages"] == [
        {"reference": "Paragraphe 1", "text": "Passage en poste avec week-ends."}
    ]
    assert documents[0]["truncated"] is False
    assert "Contenu factuel du document avenant.txt" in query
    assert "Passage en poste avec week-ends." in query


def test_document_facts_drive_the_legal_issue_instead_of_the_vague_question() -> None:
    from SYNDICAL_REASONING_ENGINE.factual_core import build_case_factual_core

    documents = SERVER.extract_uploaded_documents(
        [
            attachment(
                "projet-avenant.txt",
                (
                    "Projet d’avenant : passage d’un horaire de jour à un cycle en poste "
                    "incluant les week-ends et jours fériés."
                ).encode("utf-8"),
            )
        ]
    )
    query = SERVER.build_enriched_analysis_query(
        "Le salarié peut-il refuser le changement décrit dans le document joint ?",
        {
            "user_question": "Le salarié peut-il refuser le changement décrit dans le document joint ?",
            "uploaded_documents": documents,
        },
    )
    core = build_case_factual_core(query, "QUESTION_SALARIE")
    assert core.event_category == "WORK_SCHEDULE_CHANGE"


def test_plain_employee_wording_passer_en_poste_is_recognized_without_a_document() -> None:
    from SYNDICAL_REASONING_ENGINE.factual_core import build_case_factual_core

    core = build_case_factual_core(
        (
            "Un salarié travaille de jour. La direction veut le passer en poste "
            "avec week-ends et jours fériés et lui demande un avenant. Peut-il refuser ?"
        ),
        "QUESTION_SALARIE",
    )

    assert core.event_category == "WORK_SCHEDULE_CHANGE"


def test_upload_rejects_path_names_unsupported_formats_and_oversized_batches() -> None:
    documents = SERVER.extract_uploaded_documents(
        [attachment(r"C:\secret\note.md", b"Texte utile")]
    )
    assert documents[0]["name"] == "note.md"
    assert "C:\\" not in documents[0]["name"]
    with pytest.raises(ValueError, match="Format non accepté"):
        SERVER.extract_uploaded_documents([attachment("preuve.exe", b"danger")])
    with pytest.raises(ValueError, match="Trois pièces jointes maximum"):
        SERVER.extract_uploaded_documents(
            [attachment(f"{index}.txt", b"ok") for index in range(4)]
        )


@pytest.mark.parametrize("extension", ["jpg", "jpeg", "JPG", "JPEG"])
def test_jpeg_document_is_read_by_local_ocr(extension: str) -> None:
    image = Image.new("RGB", (1500, 220), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 48)
    draw.text(
        (30, 60),
        "Passage en poste avec week-ends et jours feries",
        fill="black",
        font=font,
    )
    stream = io.BytesIO()
    image.save(stream, format="JPEG", quality=95)

    documents = SERVER.extract_uploaded_documents(
        [attachment(f"photo-document.{extension}", stream.getvalue(), "image/jpeg")]
    )

    assert documents[0]["extension"] == f".{extension.lower()}"
    assert documents[0]["page_count"] == 1
    assert documents[0]["paragraph_count"] >= 1
    assert "Passage en poste" in documents[0]["text"]
    assert "week-ends" in documents[0]["text"]
    assert documents[0]["passages"][0]["reference"].startswith("OCR ligne")


def test_invalid_jpeg_document_is_rejected() -> None:
    with pytest.raises(ValueError, match="JPEG est invalide ou illisible"):
        SERVER.extract_uploaded_documents(
            [attachment("fausse-photo.jpg", b"not-a-jpeg", "image/jpeg")]
        )


def test_png_content_downloaded_with_a_jpeg_extension_is_normalized_and_read() -> None:
    image = Image.new("RGB", (1500, 220), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 48)
    draw.text(
        (30, 60),
        "Reglement interieur et consigne de securite",
        fill="black",
        font=font,
    )
    stream = io.BytesIO()
    image.save(stream, format="PNG")

    documents = SERVER.extract_uploaded_documents(
        [attachment("image-telephone.jpeg", stream.getvalue(), "image/jpeg")]
    )

    assert "Reglement interieur" in documents[0]["text"]
    assert "consigne de securite" in documents[0]["text"]


def test_word_agenda_preserves_direction_points_and_member_questions() -> None:
    from docx import Document
    from cse_agenda_analysis import extract_agenda_points

    document = Document()
    document.add_paragraph("COMITÉ SOCIAL ET ÉCONOMIQUE")
    document.add_paragraph("Réunion ordinaire du 12 mars 2026")
    for index in range(1, 13):
        document.add_paragraph(f"Point de la direction numéro {index}", style="List Number")
    document.add_paragraph("Questions des membres du CSE", style="List Number")
    stream = io.BytesIO()
    document.save(stream)

    documents = SERVER.extract_uploaded_documents(
        [attachment("ordre-du-jour.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")]
    )
    points = extract_agenda_points(documents[0])

    assert [point["position"] for point in points] == [str(index) for index in range(1, 14)]
    assert points[0]["title"] == "Point de la direction numéro 1"
    assert points[11]["title"] == "Point de la direction numéro 12"
    assert points[12]["title"] == "Questions des membres du CSE"


def test_interface_exposes_paste_feedback_and_local_file_picker() -> None:
    html = (SERVER_PATH.parent / "index.html").read_text(encoding="utf-8")
    javascript = (SERVER_PATH.parent / "app.js").read_text(encoding="utf-8")
    assert 'id="questionDocuments"' in html
    assert 'id="questionDocumentDropZone"' in html
    assert 'id="followUpDocumentInput"' in html
    assert ".pdf,.docx,.txt,.md,.jpg,.jpeg" in html
    assert "image/jpeg" in html
    assert 'queryInput?.addEventListener("paste"' in javascript
    assert 'questionDocumentDropZone?.addEventListener("drop"' in javascript
    assert "attachments: uploadedQuestionDocuments" in javascript
    assert "...newDocuments" in javascript
    assert "followUpDocumentInput.click()" in javascript
    assert "nexus-effective-2" in html
    upload_code = javascript[
        javascript.index("function loadQuestionDocuments") :
        javascript.index("function loadQuestionDocuments") + 2600
    ]
    assert "localStorage" not in upload_code


def test_interface_reuses_public_sources_in_the_legal_source_layers() -> None:
    javascript = (SERVER_PATH.parent / "app.js").read_text(encoding="utf-8")

    assert "function publicSourcesForLayers(publicSummary)" in javascript
    assert "publicSummary.source_extractions || publicSummary.sources || []" in javascript
    assert "renderSources(sourcesList, answer, orchestration, publicSummary)" in javascript
    assert 'source_layer: source.source_layer || "autre"' in javascript


def test_document_extraction_does_not_write_or_expose_a_local_path() -> None:
    server_source = SERVER_PATH.read_text(encoding="utf-8")
    extraction_source = server_source[
        server_source.index("def _extract_uploaded_document") :
        server_source.index("def optional_dependency_status")
    ]
    assert "write_text" not in extraction_source
    assert "write_bytes" not in extraction_source
    assert "mkstemp" not in extraction_source
    documents = SERVER.extract_uploaded_documents(
        [attachment(r"C:\Users\salarié\contrat.txt", b"Clause horaires")]
    )
    assert documents[0]["name"] == "contrat.txt"
    assert documents[0]["text"] == "Clause horaires"
    assert "C:\\" not in str(documents)


def test_uploaded_document_becomes_a_bounded_traceable_public_source() -> None:
    documents = SERVER.extract_uploaded_documents(
        [
            attachment(
                "projet-avenant.txt",
                (
                    "Objet du document.\n\nLe salarié passerait d’un horaire de jour "
                    "à un cycle en poste comprenant les week-ends et jours fériés."
                ).encode("utf-8"),
            )
        ]
    )
    result = {
        "answer": {
            "case_factual_core": {
                "event_category": "WORK_SCHEDULE_CHANGE",
                "primary_grievance_or_decision": "Passage en cycle posté",
                "facts_certain": ["Le projet inclut les week-ends."],
            }
        },
        "public_summary": {"source_extractions": []},
        "detailed_analysis": {},
    }

    projected = SERVER.add_uploaded_documents_to_public_result(
        result,
        documents,
        "Le salarié peut-il refuser le passage en poste ?",
    )
    source = projected["public_summary"]["source_extractions"][0]

    assert source["nature"] == "USER_DOCUMENT"
    assert source["availability_status"] == "PROVIDED_AND_READ"
    assert source["reference"] == "Paragraphe 2"
    assert "week-ends et jours fériés" in source["excerpt"]
    assert "ne constitue pas" in source["practical_scope"]
    assert "content_base64" not in str(projected)
    assert "C:\\" not in str(projected)


def test_health_exposes_only_boolean_connector_and_corpus_availability(monkeypatch) -> None:
    monkeypatch.setenv("CFDT_NEXUS_LEGIFRANCE_CLIENT_ID", "private-id")
    monkeypatch.setenv("CFDT_NEXUS_LEGIFRANCE_CLIENT_SECRET", "private-secret")
    monkeypatch.setenv("NEXUS_OFFICIAL_CONNECTORS_RUNTIME_ENABLED", "true")

    status = SERVER.runtime_capability_status()

    assert status["credentials"]["legifrance"] is True
    assert status["runtime"]["official_connectors"] is True
    assert "private-id" not in str(status)
    assert "private-secret" not in str(status)


def test_static_interface_is_not_served_from_an_obsolete_browser_cache() -> None:
    server_source = SERVER_PATH.read_text(encoding="utf-8")
    assert 'self.send_header("Cache-Control", "no-store, max-age=0")' in server_source
    assert 'self.send_header("Pragma", "no-cache")' in server_source
